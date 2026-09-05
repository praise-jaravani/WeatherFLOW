"""
SAWS station-data parser for the 32-station 2016-2025 block-format network.

This is Phase 1b (regression) and all future phases' data source. Replaces
the earlier Lekuba long-form parser (preserved at saws_parser_25node_lekuba.py)
which drove the 25-station 2013-2020 campaign in `results_saws_25node_phase1/`.

Network selection (locked 2026-08-08):
  - 32 stations, 2016-01-01 → 2025-12-31 (10 years, 87,672 hours)
  - All 6 SAWS variables present
  - ≤10% overall missingness across the 6 variables (Davidson (2022) criterion)
  - Independently verified by Génevieve Chikwanha, Zoleka Fuyana, and this parser
    from data/saws/audit/audit_60node.py (all three independently arrived at N=32)

Rain semantics: block-format legend says a blank cell means "no rain this
hour" (= 0), NOT missing. Only literal markers ('-', '***', '---', '=') mean
missing. This is DIFFERENT from Lekuba long-form (which uses explicit 0 and
treats blank as missing). Every rain cell here is defended by the audit
convention documented in data/saws/AUDIT.md.

Block format per sheet (one sheet per feature):
  LEGEND rows (skipped)
  HOURLY DATA :  <VarName>  -  <Month> <Year>
  <Station name>  -  Climate Number:XXXXXX  Lat:YY.YY  Lon:ZZ.ZZ  Height:HH
  DD, h01, h02, ..., h24, avg, mx, mxtm
  1  <val>  <val>  ...  (day 1, 24 hourly values in cols 1..24)
  ...
  <N>  <val>  ...       (day N)
  avg / max / tot       (summary rows to skip)
  HOURLY DATA : ...     (next block)
"""

from __future__ import annotations

import logging
import re
from calendar import monthrange
from datetime import datetime
from pathlib import Path
from typing import Iterable, List, Sequence, Tuple

import numpy as np
import pandas as pd
from openpyxl import load_workbook

log = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Canonical schema (matches the old parser's contract — downstream unchanged)
# ---------------------------------------------------------------------------
CANONICAL_VARS = ["t2m", "tp", "rh", "wd", "ws", "sp"]

# Sheet-name → canonical-var mapping. Sheet names from the 19 workbooks.
SHEET_TO_CANONICAL = {
    "Temperature":   "t2m",
    "Rain":          "tp",
    "Humidity":      "rh",
    "Wind Direction": "wd",
    "WindDirection": "wd",   # observed variant
    "Wind Speed":    "ws",
    "WindSpeed":     "ws",   # observed variant
    "Pressure":      "sp",
}

# Missing-value string markers per SAWS legend.
MISSING_MARKERS = {"-", "***", "---", "="}

# Time window (locked; matches audit + all upstream analysis)
WINDOW_START = pd.Timestamp("2016-01-01 00:00")
WINDOW_END   = pd.Timestamp("2025-12-31 23:00")

# 32-station whitelist (uppercase, stripped). Sorted by overall missingness ascending
# per audit_60node.py output on 2026-08-08. If the audit is re-run and station names
# shift (extraction variants, canonicalisation), this list needs to move with it.
STATION_WHITELIST_32 = [
    "NGQURA - COEGA",
    "MOUNT EDGECOMBE",
    "ORIBI AIRPORT",
    "ADDO ELEPHANT PARK",
    "BLOEMFONTEIN WO",
    "PORT ELIZABETH AWOS",
    "GRAAFF - REINET",
    "GRAHAMSTOWN",
    "CEDARA",
    "SOMERSET EAST",
    "DOHNE - AGR",
    "KOKSTAD",
    "BETHLEHEM WO",
    "PRETORIA UNIVERSITY PROEFPLAAS",
    "MARGATE",
    "MAKATINI RESEARCH CENTRE",
    "KNELLPOORTDAM",
    "ALIWAL-NORTH PLAATKOP",
    "PIETERMARITZBURG",
    "FAURESMITH",
    "EAST LONDON WO",
    "ROYAL NATIONAL PARK",
    "FICKSBURG",
    "CAPE ST FRANCIS",
    "LADYSMITH",
    "RIVERVIEW",
    "ELLIOT",
    "UMTHATHA WO",
    "ZUURBEKOM AWS",
    "PORT ALFRED - AIRPORT",
    "PONGOLA",
    "VAN REENEN",
]
assert len(STATION_WHITELIST_32) == 32, "whitelist must be exactly 32 stations"

WHITELIST_SET = set(STATION_WHITELIST_32)

MONTH_RE = re.compile(
    r"(January|February|March|April|May|June|July|August|September|October|November|December)\s+(\d{4})"
)
STATION_META_RE = re.compile(
    r"^(?P<name>.*?)\s+-\s+Climate\s+Number:\s*(?P<climate>[0-9 A-Za-z]+)"
    r".*?Lat:\s*(?P<lat>-?[\d.]+)"
    r".*?Lon:\s*(?P<lon>-?[\d.]+)"
    r".*?Height:\s*(?P<height>-?[\d.]+)",
    re.IGNORECASE,
)


# ---------------------------------------------------------------------------
# Cell-value coercion
# ---------------------------------------------------------------------------

def _coerce_cell(v, is_rain: bool) -> float:
    """Convert a raw cell value to float. Applies the rain-blank-zero convention.

    - Rain: blank / None / NaN → 0.0 (dry hour); markers → NaN
    - Non-rain: blank / None / NaN → NaN; markers → NaN
    - Numeric strings → float
    """
    if v is None:
        return 0.0 if is_rain else float("nan")
    if isinstance(v, float) and np.isnan(v):
        return 0.0 if is_rain else float("nan")
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if s == "":
        return 0.0 if is_rain else float("nan")
    if s in MISSING_MARKERS:
        return float("nan")
    try:
        return float(s)
    except (ValueError, TypeError):
        return float("nan")


# ---------------------------------------------------------------------------
# Workbook + block iteration
# ---------------------------------------------------------------------------

def _iter_workbook_sheets(path: Path):
    """Yield (sheet_name, rows) for every sheet in an .xlsx workbook."""
    wb = load_workbook(str(path), read_only=True, data_only=True)
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        yield sheet_name, rows
    wb.close()


def _scan_blocks(workbook_path: Path):
    """Yield (station_norm, feature_canon, ymd_hour, value, lat, lon, climate) tuples
    for every hourly cell in every block of the workbook.

    ymd_hour is a Python datetime at hour granularity.
    """
    for sheet_name, rows in _iter_workbook_sheets(workbook_path):
        feature_canon = SHEET_TO_CANONICAL.get(sheet_name.strip())
        if feature_canon is None:
            continue
        is_rain = (feature_canon == "tp")
        n_rows = len(rows)

        i = 0
        while i < n_rows:
            row = rows[i]
            if not row or row[0] is None:
                i += 1
                continue
            c0 = row[0]
            if not isinstance(c0, str) or not c0.strip().upper().startswith("HOURLY DATA"):
                i += 1
                continue

            m = MONTH_RE.search(c0)
            if not m:
                i += 1
                continue
            month = datetime.strptime(m.group(1), "%B").month
            year = int(m.group(2))
            if year < WINDOW_START.year or year > WINDOW_END.year:
                # skip blocks entirely outside our window
                i += 1
                continue

            if i + 1 >= n_rows:
                break
            meta_str = str(rows[i + 1][0]) if rows[i + 1] and rows[i + 1][0] is not None else ""
            sm = STATION_META_RE.match(meta_str)
            if not sm:
                i += 2
                continue

            station_raw = sm.group("name").strip()
            station_norm = station_raw.upper()
            if station_norm not in WHITELIST_SET:
                # not one of our 32 — skip block
                # walk past its data rows to avoid re-parsing
                j = i + 3
                while j < n_rows:
                    r = rows[j]
                    if not r or r[0] is None:
                        break
                    s = str(r[0]).strip().lower()
                    if s in ("avg", "max", "min", "tot", "med") or not s.replace(".", "").isdigit():
                        break
                    j += 1
                i = j
                continue

            climate = sm.group("climate").strip()
            lat = float(sm.group("lat"))
            lon = float(sm.group("lon"))

            # Data rows start at i+3 (i=marker, i+1=meta, i+2=header DD/h01..h24)
            data_start = i + 3
            n_days_expected = monthrange(year, month)[1]
            j = data_start
            while j < n_rows:
                r = rows[j]
                if not r or r[0] is None:
                    break
                c = r[0]
                s = str(c).strip().lower()
                if s in ("avg", "max", "min", "tot", "med") or not s.replace(".", "").isdigit():
                    break
                try:
                    day = int(float(s))
                except ValueError:
                    break
                if day < 1 or day > n_days_expected:
                    break
                for h in range(1, 25):
                    if h >= len(r):
                        break
                    val = _coerce_cell(r[h], is_rain)
                    ymd_hour = datetime(year, month, day, h - 1)
                    yield (station_norm, feature_canon, ymd_hour, val, lat, lon, climate)
                j += 1
            i = j


# ---------------------------------------------------------------------------
# Public entry — matches contract of the old parser
# ---------------------------------------------------------------------------

def parse_saws_files(raw_dir: str) -> Tuple[np.ndarray, np.ndarray, pd.DataFrame, List[str]]:
    """Parse the 60-node-campaign block-format workbooks into the pipeline contract.

    Args:
      raw_dir: directory containing .xlsx block-format workbooks. Production:
               data/saws/raw/SAWS-60-node-campaign/SAWS/extracted/
               The pipeline points here via SAWS_RAW_DIR.

    Returns:
      X          : (T=87672, N=32, D=6) float32, may contain NaN (imputation
                   downstream in preprocessing.impute)
      timestamps : (T,) datetime64[ns], hourly grid 2016-01-01 → 2025-12-31
      coords     : DataFrame [node_id, station, lat, lon, climate], N rows,
                   alphabetically ordered by station (deterministic)
      var_names  : list of D canonical names (t2m, tp, rh, wd, ws, sp)
    """
    raw = Path(raw_dir)
    if not raw.exists():
        raise FileNotFoundError(f"SAWS raw dir does not exist: {raw}")

    workbooks = sorted(raw.glob("*.xlsx")) + sorted(raw.glob("*.xls"))
    if not workbooks:
        raise FileNotFoundError(f"No workbooks in {raw}")
    log.info("Reading %d workbooks from %s", len(workbooks), raw)

    # Canonical timestamp grid
    canonical = pd.date_range(WINDOW_START, WINDOW_END, freq="h")
    T = len(canonical)
    ts_to_idx = {ts.to_pydatetime(): i for i, ts in enumerate(canonical)}

    # Station index (alphabetical order = deterministic)
    stations = sorted(STATION_WHITELIST_32)
    N = len(stations)
    stn_to_idx = {s: i for i, s in enumerate(stations)}

    # Variable index (canonical order)
    D = len(CANONICAL_VARS)
    var_to_idx = {v: i for i, v in enumerate(CANONICAL_VARS)}

    # Tensor — start all-NaN (imputation happens downstream)
    X = np.full((T, N, D), np.nan, dtype=np.float32)

    # Coord metadata per station (first-seen values used; blocks are consistent)
    coord_records: dict[str, dict] = {}

    for wb_path in workbooks:
        log.info("  parsing %s", wb_path.name)
        n_cells = 0
        for (stn, feat, ymd_hour, val, lat, lon, climate) in _scan_blocks(wb_path):
            t_idx = ts_to_idx.get(ymd_hour)
            if t_idx is None:
                continue
            n = stn_to_idx[stn]
            d = var_to_idx[feat]
            X[t_idx, n, d] = val
            n_cells += 1
            if stn not in coord_records:
                coord_records[stn] = {"lat": lat, "lon": lon, "climate": climate}
        log.info("    %d cells written", n_cells)

    # Coords DataFrame
    missing_coords = [s for s in stations if s not in coord_records]
    if missing_coords:
        raise RuntimeError(
            f"{len(missing_coords)} whitelisted stations had no data block in any workbook: "
            f"{missing_coords}"
        )
    coords = pd.DataFrame([
        {"node_id": i, "station": s,
         "lat": coord_records[s]["lat"],
         "lon": coord_records[s]["lon"],
         "climate": coord_records[s]["climate"]}
        for i, s in enumerate(stations)
    ])

    # Timestamps
    timestamps = canonical.to_numpy(dtype="datetime64[ns]")

    # Sanity — report missingness per station
    per_station_missing = np.isnan(X).sum(axis=(0, 2)) / (T * D)
    log.info("Per-station overall missingness (fraction):")
    for s, p in zip(stations, per_station_missing):
        log.info("  %-32s %6.3f", s, p)
    max_pct = float(per_station_missing.max()) * 100
    log.info("Worst station: %.2f%% overall missing (must be <=10%% per audit)", max_pct)

    log.info("Parsed: X=%s, T=%d, N=%d, D=%d", X.shape, T, N, D)
    return X, timestamps, coords, CANONICAL_VARS
